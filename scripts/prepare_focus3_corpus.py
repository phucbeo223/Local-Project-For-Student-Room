"""Add a short authentic Google Lens excerpt to a fresh Word-only trial corpus.

Reference answers are never read here. Existing 30 source entries are retained
byte for byte; raw public HTML is captured only for local provenance checks.
"""
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from urllib.request import Request, urlopen
from docx import Document
from lxml import html

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/legal_word_focus3_corpus_v10_20261006'
URL = 'https://support.google.com/websearch/answer/1325808?co=GENIE.Platform%3DDesktop&hl=vi'
SCHEMA = 'legal_word_focus3_v10_20261006'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    capture = ROOT / 'eval/provider_sources_20261005/focus3_20261006'
    capture.mkdir(parents=True, exist_ok=True)
    raw_path = capture / 'google-lens.html'
    if not raw_path.exists():
        with urlopen(Request(URL, headers={'User-Agent': 'Mozilla/5.0'}), timeout=35) as response:
            assert response.status == 200 and response.url.startswith('https://support.google.com/')
            raw_path.write_bytes(response.read())
    tree = html.fromstring(raw_path.read_bytes().decode('utf-8'))
    all_text = ' '.join(tree.text_content().split())
    segments = ['Trang web có hình ảnh đó hoặc có hình ảnh tương tự',
                'nhấp vào Tìm bằng Google Ống kính.']
    assert all(part in all_text for part in segments)
    assert sum(len(part.split()) for part in segments) <= 25
    word = OUT / 'google-lens-publisher-excerpts.docx'
    if not word.exists():
        doc = Document()
        for part in segments:
            doc.add_paragraph(part)
        doc.save(word)
    assert [p.text for p in Document(word).paragraphs] == segments
    body = '\n\n'.join(segments)
    source = dict(id='google-lens-word', title='Google Lens: kiểm tra nguồn hình ảnh — Google Search Help',
                  category='ecommerce_platform', path=word.relative_to(ROOT).as_posix(),
                  sha256=sha(word), source_url=URL, publisher='Google Search Help',
                  publication_date=None, page_kind='logical_document', pages=1, ocr_used=False,
                  extraction='Native DOCX XML; short exact publisher excerpts, no OCR/PDF',
                  source_content_kind='publisher_guidance_word_conversion',
                  original_html_sha256=sha(raw_path),
                  retrieved_at_utc=datetime.now(timezone.utc).isoformat(),
                  source_scope_warning='Hướng dẫn công cụ Google Lens, không phải điều luật. Kết quả ảnh tương tự là đầu mối để đối chiếu, không tự chứng minh tin đăng giả hoặc danh tính người cho thuê.',
                  source_policy='Only literal public publisher excerpts; no reference answers or authored guidance embedded.')
    provision = dict(provision_id='google-lens:publisher-excerpts', article=None, clause=None,
                     heading='Google Lens — tìm ảnh và trang chứa ảnh tương tự', content=body,
                     source_segments=segments, source_start=0, source_end=len(body),
                     content_sha256=hashlib.sha256(body.encode()).hexdigest(),
                     extraction_text_sha256=hashlib.sha256(body.encode()).hexdigest(),
                     points=[], exceptions=[], cross_references=[], kind='publisher_guidance',
                     article_context='', page_from=1, page_to=1)
    target = OUT / 'google-lens-word.json'
    save(target, dict(source=source, provisions=[provision]))
    base_path = ROOT / 'docs/legal_word_supplement_corpus_v8_20261005/manifest.json'
    base = json.loads(base_path.read_text(encoding='utf-8'))
    assert all(sha(ROOT / entry['file']) == entry['sha256'] for entry in base['documents'])
    entry = dict(id=source['id'], file=target.relative_to(ROOT).as_posix(),
                 sha256=sha(target), category=source['category'], provisions=1)
    save(OUT / 'manifest.json', dict(schema=SCHEMA, documents=[*base['documents'], entry],
                                   base_manifest_sha256=sha(base_path), base_documents_unchanged=True,
                                   policy='Fresh isolated trial; old v7 and v8 stores retained. Word-only exact sources, no OCR/PDF or reference embedding.'))
    print(json.dumps(dict(schema=SCHEMA, documents=31, added_source=URL,
                          quoted_words=sum(len(p.split()) for p in segments),
                          original_30_entries_unchanged=True), ensure_ascii=False))


if __name__ == '__main__':
    main()
