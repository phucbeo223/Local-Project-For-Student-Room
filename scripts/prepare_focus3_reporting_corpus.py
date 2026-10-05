"""Add a literal publisher reporting-evidence excerpt; retain the v10 trial."""
import hashlib
import json
from datetime import datetime, timezone
from pathlib import Path
from urllib.request import Request, urlopen
from docx import Document
from lxml import html

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / 'docs/legal_word_focus3_corpus_v11_20261006'
URL = 'https://trogiup.chotot.com/nguoi-mua/bi-quyet-mua-bat-dong-san-an-toan-voi-nha-tot/'


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def save(path, value):
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')


def main():
    OUT.mkdir(parents=True, exist_ok=True)
    raw_path = ROOT / 'eval/provider_sources_20261005/focus3_20261006/nhatot-report-proof.html'
    if not raw_path.exists():
        with urlopen(Request(URL, headers={'User-Agent': 'Mozilla/5.0'}), timeout=35) as response:
            assert response.status == 200 and response.url.startswith('https://trogiup.chotot.com/')
            raw_path.write_bytes(response.read())
    original = ' '.join(html.fromstring(raw_path.read_bytes().decode('utf-8')).text_content().split())
    excerpt = 'Thông tin (hình ảnh) trao đổi mua bán của bạn với người bán thể hiện hành vi vi phạm (nếu có).'
    assert excerpt in original and len(excerpt.split()) <= 25
    word = OUT / 'nhatot-report-proof.docx'
    if not word.exists():
        doc = Document(); doc.add_paragraph(excerpt); doc.save(word)
    assert [p.text for p in Document(word).paragraphs] == [excerpt]
    source = dict(id='nhatot-report-proof-word', title='Nhà Tốt: hình ảnh trao đổi hỗ trợ phản ánh tin vi phạm',
                  category='ecommerce_platform', path=word.relative_to(ROOT).as_posix(), sha256=sha(word),
                  source_url=URL, publisher='Trợ Giúp Chợ Tốt / Nhà Tốt', publication_date=None,
                  page_kind='logical_document', pages=1, ocr_used=False,
                  extraction='Native DOCX XML; exact short publisher excerpt, no OCR/PDF',
                  source_content_kind='publisher_guidance_word_conversion',
                  original_html_sha256=sha(raw_path), retrieved_at_utc=datetime.now(timezone.utc).isoformat(),
                  source_scope_warning='Hướng dẫn của Nhà Tốt về hình ảnh trao đổi giúp làm rõ phản ánh tin/người bán vi phạm; không phải điều luật hoặc thời hạn xử lý. Trang nguồn có bối cảnh mua bất động sản; chỉ dùng phần phản ánh tin, không suy sang thủ tục mua bán hoặc quyền trong hợp đồng thuê.',
                  source_policy='Only literal public publisher excerpt; no reference answers or authored legal rules embedded.')
    provision = dict(provision_id='nhatot:report-proof', article=None, clause=None,
                     heading='Hình ảnh trao đổi khi phản ánh tin đăng/người bán vi phạm', content=excerpt,
                     source_segments=[excerpt], source_start=0, source_end=len(excerpt),
                     content_sha256=hashlib.sha256(excerpt.encode()).hexdigest(),
                     extraction_text_sha256=hashlib.sha256(excerpt.encode()).hexdigest(),
                     points=[], exceptions=[], cross_references=[], kind='publisher_guidance',
                     article_context='', page_from=1, page_to=1)
    target = OUT / 'nhatot-report-proof-word.json'
    save(target, dict(source=source, provisions=[provision]))
    base_path = ROOT / 'docs/legal_word_focus3_corpus_v10_20261006/manifest.json'
    base = json.loads(base_path.read_text(encoding='utf-8'))
    assert all(sha(ROOT / entry['file']) == entry['sha256'] for entry in base['documents'])
    entry = dict(id=source['id'], file=target.relative_to(ROOT).as_posix(), sha256=sha(target),
                 category=source['category'], provisions=1)
    save(OUT / 'manifest.json', dict(schema='legal_word_focus3_v11_20261006', documents=[*base['documents'], entry],
                                   base_manifest_sha256=sha(base_path), base_documents_unchanged=True,
                                   policy='Fresh isolated trial; original active store and v8/v10 retained. Word-only authentic sources; references never embedded.'))
    print(json.dumps(dict(documents=32, added_source=URL, quoted_words=len(excerpt.split()),
                          original_31_entries_unchanged=True), ensure_ascii=False))


if __name__ == '__main__':
    main()
