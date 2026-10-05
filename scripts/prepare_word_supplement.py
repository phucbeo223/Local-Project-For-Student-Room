"""Literal public sources + unchanged v7 corpus; never read reference answers."""
from pathlib import Path
import hashlib
import json
import sys
import importlib.util
from docx import Document
from lxml import html

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('supplement_provisions',ROOT/'apps/api/app/room_service/legal_knowledge/provisions.py')
module=importlib.util.module_from_spec(spec)
sys.modules[spec.name]=module
spec.loader.exec_module(module)
structure_provisions,provision_json=module.structure_provisions,module.provision_json

OUT=ROOT/'docs/legal_word_supplement_corpus_v8_20261005'
CAP=ROOT/'eval/provider_sources_20261005/word_supplement_v8'
SCHEMA='legal_word_supplement_v8_20261005'
SELECTIONS={
 'S01':('housing_contract',list(range(78,90))+list(range(91,94))+list(range(95,105))),
 'S02':('housing_contract',[23,24,25,30]),
 'S03':('electricity',[61,63,158,160]),
 'S04':('ecommerce_platform',list(range(75,82))+[86]),
 'S05':('water_cantho',list(range(5))),
 'S06':('water_cantho',[45,46,50,51,52,53,54,55,56,60]),
 'S08':('ecommerce_platform',list(range(93,110))),
 'S09':('ecommerce_platform',list(range(75,80))+list(range(83,93))),
 'S10':('ecommerce_platform',list(range(370,395))+list(range(401,413))),
 'S11':('criminal_law',[12,19,20,22]),
 'S12':('criminal_law',[12,14,15,16,17]),
 'S14':('fire_safety',list(range(2,16))+list(range(18,27))),
}
LAWS={
 'S15':('Luat-122-2025-QH15.docx',[3,11,15,17,18,21,40,41]),
 'S16':('Nghi-dinh-248-2026-ND-CP.docx',[2,4,17,18,19,52,53]),
}
def digest(path): return hashlib.sha256(path.read_bytes()).hexdigest()
def save(path,value): path.write_text(json.dumps(value,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')

def build():
 OUT.mkdir(parents=True,exist_ok=True)
 base_path=ROOT/'docs/legal_word_corpus_v7_20261005/manifest.json'
 base=json.loads(base_path.read_text(encoding='utf-8'))
 entries=list(base['documents'])
 sources=json.loads((ROOT/'docs/review_sources_20261005/sources.json').read_text(encoding='utf-8'))['sources']
 captures={r['id']:r for r in json.loads((CAP/'capture.json').read_text(encoding='utf-8'))}
 originals={r['id']:r for r in json.loads((ROOT/'docs/review_sources_20261005/originals/provenance.json').read_text(encoding='utf-8'))}
 fidelity=[]
 for s in sources:
  key=s['id']
  if key not in SELECTIONS and key not in LAWS: continue
  literal=[]
  if key in SELECTIONS:
   category,indices=SELECTIONS[key]
   assert digest(CAP/(key+'.html'))==captures[key]['sha256']
   blocks=json.loads((CAP/(key+'.blocks.json')).read_text(encoding='utf-8'))
   if key=='S14':
    tree=html.fromstring((CAP/(key+'.html')).read_text(encoding='utf-8'))
    blocks=[' '.join(e.text_content().split()) for e in tree.xpath('//div[@id="articleContent"]/div') if e.text_content().strip()]
   assert all(i<len(blocks) for i in indices)
   literal=[blocks[i] for i in indices]
   assert all(literal) and not any('\ufffd' in p or 'Ã¡' in p or 'á»' in p for p in literal), key
   word=OUT/(key+'-publisher-text.docx')
   if not word.exists():
    doc=Document()
    for p in literal: doc.add_paragraph(p)
    doc.save(word)
   assert [p.text for p in Document(word).paragraphs]==literal
   body='\n\n'.join(literal)
   category_warning='Hướng dẫn/khuyến cáo của '+s['publisher']+'; không phải điều luật. '+s['limitations']
   kind='publisher_guidance_word_conversion'
   # Keep related instructions together rather than embedding lone headings.
   groups=[]; current=[]
   for i,p in zip(indices,literal):
    if current and sum(len(v) for _,v in current)+len(p)>2200:
     groups.append(current);current=[]
    current.append((i,p))
   if current: groups.append(current)
   provisions=[]; offset=0
   for group in groups:
    i=group[0][0];p='\n\n'.join(v for _,v in group)
    start=body.index(p,offset); offset=start+len(p)
    h=hashlib.sha256(p.encode()).hexdigest()
    provisions.append(dict(provision_id=key+':web-block:'+str(i),article=None,clause=None,
      heading=s['title']+' | Đoạn nguyên văn '+str(i),content=p,source_start=start,source_end=offset,
      points=[],exceptions=[],cross_references=[],kind='publisher_guidance',article_context='',
      content_sha256=h,page_from=1,page_to=1,original_block_indices=[n for n,_ in group]))
  else:
   filename,articles=LAWS[key]
   word=ROOT/'docs/review_sources_20261005/originals'/filename
   assert digest(word)==originals[key]['sha256'] and originals[key]['tls_verified']
   body='\n\n'.join(p.text for p in Document(word).paragraphs if p.text.strip())
   parts=[p for p in structure_provisions(body,digest(word)) if p.article in set(map(str,articles))]
   assert set(p.article for p in parts)==set(map(str,articles)),key
   provisions=[]
   for p in parts:
    assert p.content in body
    v=provision_json(p);v.update(page_from=1,page_to=1);provisions.append(v)
   category='ecommerce_platform'; kind='published_word'
   category_warning=('Văn bản TMĐT có hiệu lực 01/07/2026; phải đọc đúng loại nền tảng và điều khoản chuyển tiếp. '
      +('Điều 52 khoản 2: quy định xác thực điện tử áp dụng từ 01/01/2027, chưa áp dụng ở ngày 05/10/2026.' if key=='S16' else 'Điều 41: hồ sơ nền tảng đã xác nhận trước 01/07/2026 được tiếp tục đến 30/06/2027.'))
  text_sha=hashlib.sha256(body.encode()).hexdigest()
  for v in provisions:
   assert v['content'] in body
   v['extraction_text_sha256']=text_sha
  source=dict(id='supplement-'+key.lower()+'-word',title=s['title']+' — '+s['publisher'],category=category,
   path=word.relative_to(ROOT).as_posix(),sha256=digest(word),source_url=s['url'],publisher=s['publisher'],
   publication_date=s['publication_date'],page_kind='logical_document',pages=1,ocr_used=False,
   extraction='Native DOCX XML; literal publisher text, no OCR or PDF',source_content_kind=kind,
   source_scope_warning=category_warning.replace(' Đợt này chỉ đọc hướng dẫn, không gửi báo cáo.',''),not_exhaustive=True,
   source_policy='Literal publisher text; never assistant summary, reference answer, form or drafted agreement.')
  if key in SELECTIONS: source.update(original_html_sha256=captures[key]['sha256'],original_block_indices=indices,conversion='Whitespace normalization only; selected exact DOM blocks converted to DOCX',retrieved_at_utc=captures[key]['retrieved_at_utc'])
  else: source.update(original_sha256=originals[key]['sha256'],download_url=originals[key]['download_url'],retrieved_at_utc=originals[key]['retrieved_at_utc'])
  target=OUT/(source['id']+'.json');save(target,dict(source=source,provisions=provisions))
  entries.append(dict(id=source['id'],file=target.relative_to(ROOT).as_posix(),sha256=digest(target),category=category,provisions=len(provisions)))
  fidelity.append(dict(id=key,word_sha256=digest(word),literal_word_roundtrip=True,provisions=len(provisions),selected_characters=sum(len(p['content']) for p in provisions)))
 for e in base['documents']: assert digest(ROOT/e['file'])==e['sha256']
 manifest=dict(schema=SCHEMA,documents=entries,base_manifest_sha256=digest(base_path),base_documents_unchanged=True,
   policy='Old 16 Word sources plus literal selected public sources in native/converted Word; no OCR/PDF, no authored summaries or reference answers; isolated experiment, current release retained.',
   approval='Human requested embedding reviewed supplemental sources with previous documents into a new store, keeping active store.',
   excluded_supplements=[dict(id='S07',reason='Article content supplied as images; no readable native article text verified, OCR explicitly excluded. Footer price table is not the article.'),dict(id='S13',reason='Publisher HTTP 403 blocks reproducible original capture; web reading available but no native capture admitted to this release.')])
 save(OUT/'manifest.json',manifest);save(OUT/'fidelity.json',fidelity)
 print(json.dumps(dict(documents=len(entries),added=len(fidelity),provisions=sum(e['provisions'] for e in entries),manifest_sha256=digest(OUT/'manifest.json')),ensure_ascii=False))
if __name__=='__main__': build()
